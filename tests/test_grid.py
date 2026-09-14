"""원본 형식 판별과 격자 읽기. 합성 통합문서로만 확인한다."""

import io
import re
import zipfile
from datetime import datetime, time
from pathlib import Path

import openpyxl
import pytest
import xlwt

from deliciousmap.grid import UnreadableOriginal, UnsupportedFormat, read_tables
from tests import hwpx, pdf
from tests.gwangju import bundle, workbook


def test_legacy_workbook_keeps_date_cells_and_skips_empty_sheets(tmp_path: Path) -> None:
    path = tmp_path / "원본.bin"
    path.write_bytes(
        workbook([("제목",), ("일시", "금액"), (datetime(2026, 1, 5, 12, 7), 62000.0)], [])
    )
    (table,) = read_tables(path)
    assert (table.name, table.label) == ("sheet1", "시트1")
    assert table.rows[2] == (datetime(2026, 1, 5, 12, 7), 62000.0)


def test_xlsx_is_read_by_content_and_trims_formatted_empty_ranges(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "기관운영"
    sheet.append(["□ 합성과 업무추진비"])
    sheet.append(["사용자", "사용일시", "시각", "금액"])
    sheet.append(["합성과장", datetime(2026, 2, 3), time(12, 4), 93000])
    # 서식만 있는 먼 칸은 격자에 넣지 않는다.
    sheet.cell(row=5000, column=60).number_format = "0"
    book.create_sheet("빈 시트")
    path = tmp_path / "집행내역.xls"
    book.save(path)
    (table,) = read_tables(path)
    assert table.label == "기관운영"
    assert table.rows == (
        ("□ 합성과 업무추진비",),
        ("사용자", "사용일시", "시각", "금액"),
        ("합성과장", datetime(2026, 2, 3), "12:04", 93000.0),
    )


def test_vertically_merged_cells_are_marked_without_changing_the_grid(tmp_path: Path) -> None:
    """격자는 병합을 담지 않는다. 이어짐 칸은 빈 값 그대로 두고 병합만 따로 싣는다.

    값을 채워 넣으면 헤더 서명이 달라져 공통 헤더 캐시가 빗나간다(#113).
    """
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["일자", "사용장소", "사용금액"])
    sheet.append(["2026-03-17", "합성 식당", 62000])
    sheet.append([None, "합성 찻집", 27000])
    sheet.append([None, "합성 국밥", 15000])
    sheet.merge_cells("A2:A3")
    path = tmp_path / "집행내역.xlsx"
    book.save(path)
    (table,) = read_tables(path)
    assert table.rows[2] == ("", "합성 찻집", 27000.0)
    assert table.value(3, 0) == "2026-03-17"
    # 병합이 아닌 빈 칸은 그대로 빈 값이다.
    assert table.value(4, 0) == ""


def test_legacy_workbook_marks_vertically_merged_cells(tmp_path: Path) -> None:
    book = xlwt.Workbook()
    sheet = book.add_sheet("시트1")
    for column, label in enumerate(("일자", "사용장소", "사용금액")):
        sheet.write(0, column, label)
    sheet.write_merge(1, 2, 0, 0, "2026-03-17")
    spent = [("합성 식당", 62000.0), ("합성 찻집", 27000.0)]
    for row, (place, amount) in enumerate(spent, start=1):
        sheet.write(row, 1, place)
        sheet.write(row, 2, amount)
    sheet.write(3, 1, "합성 국밥")
    sheet.write(3, 2, 15000.0)
    stream = io.BytesIO()
    book.save(stream)
    path = tmp_path / "집행내역.xls"
    path.write_bytes(stream.getvalue())
    (table,) = read_tables(path)
    assert table.rows[2] == ("", "합성 찻집", 27000.0)
    assert table.value(3, 0) == "2026-03-17"
    assert table.value(4, 0) == ""


def strict(path: Path) -> bytes:
    """openpyxl이 쓴 통합문서를 ISO Strict 이름공간으로 바꾼다. 실제 원본 두 개가 이 형식이다."""
    source = zipfile.ZipFile(path)
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as target:
        for name in source.namelist():
            data = source.read(name)
            for transitional, iso in (
                (b"http://schemas.openxmlformats.org/spreadsheetml/2006/main",
                 b"http://purl.oclc.org/ooxml/spreadsheetml/main"),
                (b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
                 b"http://purl.oclc.org/ooxml/officeDocument/relationships"),
            ):  # fmt: skip
                data = data.replace(transitional, iso)
            target.writestr(name, data)
    return stream.getvalue()


def test_strict_ooxml_workbook_is_read_like_a_transitional_one(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    book.active.append(["사용일시", "금액"])
    book.active.append(["2026-01-05", 62000])
    book.save(tmp_path / "보통.xlsx")
    path = tmp_path / "엄격.xlsx"
    path.write_bytes(strict(tmp_path / "보통.xlsx"))
    (table,) = read_tables(path)
    assert table.rows == (("사용일시", "금액"), ("2026-01-05", 62000.0))


def test_workbook_without_readable_sheets_is_not_reported_as_empty(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    book.save(tmp_path / "보통.xlsx")
    source = zipfile.ZipFile(tmp_path / "보통.xlsx")
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as target:
        for name in source.namelist():
            data = source.read(name)
            if name == "xl/workbook.xml":
                data = re.sub(rb"<sheets>.*</sheets>", b"<sheets/>", data)
            target.writestr(name, data)
    path = tmp_path / "시트 없음.xlsx"
    path.write_bytes(stream.getvalue())
    with pytest.raises(UnreadableOriginal):
        read_tables(path)


@pytest.mark.parametrize(
    "content",
    [
        b"HWP Document File",
        b"PK\x03\x04not-a-workbook",
        # 게시판의 첨부 묶음처럼 통합문서도 HWPX 본문도 없는 ZIP.
        bundle(("첨부/집행내역.txt", "붙임과 같이 게시합니다.".encode())),
    ],
)
def test_other_formats_are_unsupported(tmp_path: Path, content: bytes) -> None:
    path = tmp_path / "원본.xls"
    path.write_bytes(content)
    with pytest.raises(UnsupportedFormat):
        read_tables(path)


def hwp(stream: str) -> bytes:
    """통합문서가 없는 OLE2 컨테이너. HWP 5.0이 `.hwpx`와 달리 걸리는 자리를 그대로 짚는다."""
    book = xlwt.Workbook()
    book.add_sheet("시트1").write(0, 0, "합성")
    saved = io.BytesIO()
    book.save(saved)
    return saved.getvalue().replace("Workbook".encode("utf-16-le"), stream.encode("utf-16-le"))


def test_hwp_stays_an_unsupported_format(tmp_path: Path) -> None:
    """HWP 5.0은 ZIP이 아니라 OLE2다. 이번 구현은 `.hwpx`만 읽고 `.hwp`는 그대로 미해결이다."""
    path = tmp_path / "집행내역.hwp"
    path.write_bytes(hwp("BodyText"))
    with pytest.raises(UnsupportedFormat):
        read_tables(path)


def test_hwpx_table_is_read_with_the_heading_before_it_as_the_label(tmp_path: Path) -> None:
    path = tmp_path / "집행내역.hwpx"
    path.write_bytes(
        hwpx.document(
            "2026. 1분기 업무추진비 집행내역(합성과)",
            [
                ("사용자", "일시", "장소", "금액"),
                ("합성과장", "2026.01.02.12:30", "합성식당", "50,400"),
            ],
        )
    )
    (table,) = read_tables(path)
    assert (table.name, table.label) == ("table1", "2026. 1분기 업무추진비 집행내역(합성과)")
    assert table.rows == (
        ("사용자", "일시", "장소", "금액"),
        ("합성과장", "2026.01.02.12:30", "합성식당", "50,400"),
    )


def test_hwpx_joins_the_paragraphs_one_cell_is_split_into(tmp_path: Path) -> None:
    """실제 원본은 칸 안에서 줄을 나눠 `결제`·`방법`을 따로 적는다. 한 칸은 한 값이다."""
    path = tmp_path / "집행내역.hwpx"
    path.write_bytes(
        hwpx.document(
            [
                ("일시", hwpx.cell("결제", "방법")),
                (hwpx.cell("2026. 4. 8.", "12:00"), hwpx.cell("신용", "카드")),
            ]
        )
    )
    (table,) = read_tables(path)
    assert table.rows == (("일시", "결제방법"), ("2026. 4. 8.12:00", "신용카드"))


def test_hwpx_merged_cells_keep_the_original_row_and_column_numbers(tmp_path: Path) -> None:
    """합계 행이 `colSpan`으로 합쳐져 있어도 금액은 제 열에 남고, 덮인 자리는 비운다."""
    path = tmp_path / "집행내역.hwpx"
    path.write_bytes(
        hwpx.document(
            [
                ("사용자", "일시", "장소", "금액"),
                (hwpx.cell("합 계", columns=3), "639,100"),
                (hwpx.cell("합성과장", rows=2), "2026.01.02.12:30", "합성식당", "50,400"),
                ("2026.01.14.12:00", "합성찻집", "588,700"),
            ]
        )
    )
    (table,) = read_tables(path)
    assert table.cell(2, 0) == "합 계"
    assert table.cell(2, 3) == "639,100"
    assert table.rows[1] == ("합 계", "", "", "639,100")
    assert table.rows[3] == ("", "2026.01.14.12:00", "합성찻집", "588,700")


def test_hwpx_marks_cells_that_are_merged_into_the_one_above(tmp_path: Path) -> None:
    """세로 병합은 `<hp:cellSpan>`이 밝힌다. 통합문서·PDF와 같이 격자는 그대로 두고 따로 싣는다."""
    path = tmp_path / "집행내역.hwpx"
    path.write_bytes(
        hwpx.document(
            [
                ("일자", "사용장소", "사용금액"),
                (hwpx.cell("2026-03-17", rows=2), "합성 식당", "62,000"),
                ("합성 찻집", "27,000"),
                ("", "합성 국밥", "15,000"),
            ]
        )
    )
    (table,) = read_tables(path)
    assert table.rows[2] == ("", "합성 찻집", "27,000")
    assert table.value(3, 0) == "2026-03-17"
    # 병합이 아닌 빈 칸은 그대로 빈 값이다.
    assert table.value(4, 0) == ""


def test_hwpx_numbers_the_tables_and_labels_each_with_its_own_heading(tmp_path: Path) -> None:
    path = tmp_path / "집행내역.hwpx"
    path.write_bytes(
        hwpx.document(
            "□ 1분기",
            [("일시", "장소"), ("2026.01.02.12:30", "합성식당")],
            "□ 2분기",
            [("일시", "장소"), ("2026.04.08.12:00", "합성찻집")],
        )
    )
    first, second = read_tables(path)
    assert [(table.name, table.label) for table in (first, second)] == [
        ("table1", "□ 1분기"),
        ("table2", "□ 2분기"),
    ]
    assert second.rows[1] == ("2026.04.08.12:00", "합성찻집")


def test_hwpx_numbers_the_tables_of_every_body_in_document_order(tmp_path: Path) -> None:
    """본문이 여럿인 원본. 구역 번호는 글자가 아니라 수로 이어지므로 10은 2보다 뒤다."""
    path = tmp_path / "집행내역.hwpx"
    later = hwpx.section("□ 뒤 구역", [("일시",), ("2026.05.11.12:00",)])
    earlier = hwpx.section("□ 앞 구역", [("일시",), ("2026.03.02.12:00",)])
    path.write_bytes(
        hwpx.archive(("Contents/section10.xml", later), ("Contents/section2.xml", earlier))
    )
    first, second = read_tables(path)
    assert [(table.name, table.label) for table in (first, second)] == [
        ("table1", "□ 앞 구역"),
        ("table2", "□ 뒤 구역"),
    ]


def test_hwpx_without_a_table_is_not_reported_as_empty(tmp_path: Path) -> None:
    """표 밖 문단은 표로 만들지 않는다. 표 0개를 집행 없음으로 바꾸지 않는다."""
    path = tmp_path / "안내문.hwpx"
    path.write_bytes(hwpx.document("붙임과 같이 게시합니다.", "끝."))
    assert read_tables(path) == ()


def test_broken_hwpx_body_is_unreadable_rather_than_an_unsupported_format(tmp_path: Path) -> None:
    path = tmp_path / "깨진.hwpx"
    path.write_bytes(hwpx.document(body="<hs:sec><hp:p>"))
    with pytest.raises(UnreadableOriginal):
        read_tables(path)


def test_hwpx_cells_pointing_at_one_position_are_unreadable(tmp_path: Path) -> None:
    """자리 표기가 어긋나 두 칸이 한 자리를 가리킨다. 나중 칸으로 덮어써 앞 칸을 버리지 않는다."""
    path = tmp_path / "겹친 자리.hwpx"
    cell = (
        "<hp:tc><hp:subList><hp:p><hp:run><hp:t>합성식당</hp:t></hp:run></hp:p></hp:subList>"
        '<hp:cellAddr colAddr="0" rowAddr="0"/><hp:cellSpan colSpan="1" rowSpan="1"/></hp:tc>'
    )
    path.write_bytes(
        hwpx.document(
            body='<?xml version="1.0" encoding="UTF-8" standalone="yes" ?>'
            f"<hs:sec {hwpx.NAMESPACES}><hp:p><hp:run>"
            f'<hp:tbl rowCnt="1" colCnt="2"><hp:tr>{cell}{cell}</hp:tr></hp:tbl>'
            "</hp:run></hp:p></hs:sec>"
        )
    )
    with pytest.raises(UnreadableOriginal):
        read_tables(path)


def test_hwpx_table_too_wide_to_lay_out_is_unreadable(tmp_path: Path) -> None:
    """병합 표기가 깨져 격자를 감당할 수 없는 표. 자리를 짐작해 줄이지 않고 미해결로 남긴다."""
    path = tmp_path / "깨진 병합.hwpx"
    path.write_bytes(
        hwpx.document(
            body='<?xml version="1.0" encoding="UTF-8" standalone="yes" ?>'
            f"<hs:sec {hwpx.NAMESPACES}><hp:p><hp:run>"
            '<hp:tbl rowCnt="1" colCnt="1"><hp:tr><hp:tc><hp:subList><hp:p><hp:run>'
            "<hp:t>합 계</hp:t></hp:run></hp:p></hp:subList>"
            '<hp:cellAddr colAddr="0" rowAddr="0"/><hp:cellSpan colSpan="9999999" rowSpan="1"/>'
            "</hp:tc></hp:tr></hp:tbl></hp:run></hp:p></hs:sec>"
        )
    )
    with pytest.raises(UnreadableOriginal):
        read_tables(path)


def test_pdf_table_is_read_with_the_page_as_the_label(tmp_path: Path) -> None:
    path = tmp_path / "집행내역.pdf"
    path.write_bytes(
        pdf.document(
            [
                ("사용자", "사용일시", "사용장소", "사용금액"),
                ("", "계", "", "155,000"),
                ("합성과", "2026-01-07 12:22", "합성식당", "93,000"),
                ("합성과", "2026-01-08 12:22", "합성찻집", "62,000"),
            ]
        )
    )
    (table,) = read_tables(path)
    assert (table.name, table.label) == ("table1", "1쪽")
    assert table.rows[0] == ("사용자", "사용일시", "사용장소", "사용금액")
    assert table.rows[2] == ("합성과", "2026-01-07 12:22", "합성식당", "93,000")


def test_pdf_numbers_the_tables_and_names_the_page_each_came_from(tmp_path: Path) -> None:
    path = tmp_path / "집행내역.pdf"
    path.write_bytes(
        pdf.document(
            [("사용일시", "사용장소"), ("2026-01-07", "합성식당")],
            [("사용일시", "사용장소"), ("2026-02-09", "합성찻집")],
        )
    )
    first, second = read_tables(path)
    assert [(table.name, table.label) for table in (first, second)] == [
        ("table1", "1쪽"),
        ("table2", "2쪽"),
    ]
    assert second.rows[1] == ("2026-02-09", "합성찻집")


def test_pdf_marks_cells_that_are_merged_into_the_one_above(tmp_path: Path) -> None:
    """병합된 자리에는 칸 자체가 없다(남구 실측). 그 자리를 병합으로 싣고 격자는 그대로 둔다."""
    path = tmp_path / "집행내역.pdf"
    path.write_bytes(
        pdf.document(
            [
                ("일자", "사용장소", "사용금액"),
                ("2026-03-17", "합성 식당", "62,000"),
                (None, "합성 찻집", "27,000"),
                ("", "합성 국밥", "15,000"),
            ]
        )
    )
    (table,) = read_tables(path)
    assert table.rows[2] == ("", "합성 찻집", "27,000")
    assert table.value(3, 0) == "2026-03-17"
    # 괘선이 있는 빈 칸은 병합이 아니다.
    assert table.value(4, 0) == ""


def test_pdf_without_a_readable_table_is_not_reported_as_empty(tmp_path: Path) -> None:
    """괘선이 없어 칸을 가를 수 없는 PDF. 표 0개를 집행 없음으로 바꾸지 않는다."""
    path = tmp_path / "안내문.pdf"
    path.write_bytes(pdf.document([("붙임과 같이 게시합니다.",)], ruled=False))
    assert read_tables(path) == ()


def test_broken_pdf_is_unreadable_rather_than_an_unsupported_format(tmp_path: Path) -> None:
    path = tmp_path / "깨진.pdf"
    path.write_bytes(b"%PDF-1.7 synthetic")
    with pytest.raises(UnreadableOriginal):
        read_tables(path)


def test_html_tables_are_read_with_merges_like_the_other_formats(tmp_path: Path) -> None:
    """쪽의 `<table>`마다 표 하나다. 캡션이 이름표이고, 병합 값은 왼쪽 위 칸에만 둔다."""
    page = (
        "<!DOCTYPE html><html><head><script>var cell = '<td>x</td>';</script></head><body>"
        "<h2>2026-03-17 합성 집행내역</h2>"
        "<table><caption>합성 상세</caption>"
        "<tr><th>일자</th><th colspan=2>사용 내역</th></tr>"
        "<tr><td rowspan=2>2026-03-17</td><td>합성 식당</td><td>62,000</td></tr>"
        "<tr><td>합성 찻집<br/>2층</td><td>27,000</td></tr>"
        "<tr><td></td><td>합성 국밥</td><td>15,000</td></tr>"
        "</table>"
        "<p>안내</p><table><tr><td>목록</td></tr></table></body></html>"
    )
    path = tmp_path / "20260317-1.html"
    path.write_text(page, encoding="utf-8")
    first, second = read_tables(path)
    assert (first.name, first.label) == ("table1", "합성 상세")
    # 가로 병합이 덮은 자리는 빈 값이고, 행 끝의 빈 칸은 다른 형식처럼 잘린다.
    assert first.rows == (
        ("일자", "사용 내역"),
        ("2026-03-17", "합성 식당", "62,000"),
        ("", "합성 찻집 2층", "27,000"),
        ("", "합성 국밥", "15,000"),
    )
    assert first.value(3, 0) == "2026-03-17"
    # 병합이 아닌 빈 칸은 그대로 빈 값이다.
    assert first.value(4, 0) == ""
    # 캡션이 없으면 표 바로 앞의 문단이 이름표다.
    assert (second.name, second.label) == ("table2", "안내")


def test_an_html_page_outside_the_measured_encoding_is_unreadable(tmp_path: Path) -> None:
    path = tmp_path / "1-1.html"
    path.write_bytes("<html><table><tr><td>합성</td></tr></table></html>".encode("euc-kr"))
    with pytest.raises(UnreadableOriginal):
        read_tables(path)


def test_a_repeated_label_on_an_empty_element_does_not_hide_the_rest_of_the_table(
    tmp_path: Path,
) -> None:
    # 닫는 태그가 없는 요소(`<br>`·`<img>`)에 열 이름 표시가 붙어도 뒤의 칸을 숨기지 않는다.
    page = (
        "<html><body><table><tr><th>장소</th><th>금액(원)</th></tr>"
        '<tr><td><img class="add-head" alt="장소">합성 식당</td>'
        '<td><br class="add-head"><span class="add-head">금액(원)</span>62,000</td></tr>'
        "</table></body></html>"
    )
    path = tmp_path / "1-1.html"
    path.write_text(page, encoding="utf-8")
    (table,) = read_tables(path)
    assert table.rows == (("장소", "금액(원)"), ("합성 식당", "62,000"))


def test_a_spreadsheetml_workbook_is_not_read_as_an_html_page(tmp_path: Path) -> None:
    # SpreadsheetML도 `<Table>`을 담지만 HTML 표 쪽이 아니다. 읽지 않는 형식으로 그대로 둔다.
    path = tmp_path / "집행내역.xls"
    path.write_bytes(
        b'<?xml version="1.0"?>\n<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet">'
        b"<Worksheet><Table><Row><Cell><Data>1</Data></Cell></Row></Table></Worksheet></Workbook>"
    )
    with pytest.raises(UnsupportedFormat):
        read_tables(path)
