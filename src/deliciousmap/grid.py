"""원본을 표 격자로 읽는다. 형식 판별은 확장자가 아니라 파일 앞 바이트로 한다."""

import io
import re
import warnings
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, time
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
import pdfplumber
import xlrd
from pdfplumber.table import Table as RuledTable

# 셀 값. 엑셀의 날짜 셀만 datetime이고 숫자는 float, 나머지는 앞뒤 공백을 둔 문자열이다.
Cell = str | float | datetime
# 형식별 읽기가 내는 표 하나. `Table.label`이 될 이름표, 자르기 전의 행들, 세로 병합이다.
Block = tuple[str, list[tuple[Cell, ...]], tuple["Span", ...]]

OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ZIP = b"PK\x03\x04"
PDF = b"%PDF-"
# 연속 빈 행이 이만큼 이어지면 시트의 끝으로 본다.
MAX_BLANK_RUN = 1_000
# 표 하나를 펼칠 수 있는 칸 수의 상한. 자리·병합 표기가 깨진 HWPX가 격자를 키우지 못하게 한다.
MAX_TABLE_CELLS = 1_000_000
# 통합문서 XML의 이름공간. `_transitional`이 Strict를 이 이름공간으로 바꾼 뒤에 읽는다.
MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
DOCUMENT = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
WORKBOOK_RELS = "xl/_rels/workbook.xml.rels"
# 괘선 좌표가 소수점 아래에서 어긋나는 것을 견디는 폭(포인트).
RULE_TOLERANCE = 1.0
# ISO/IEC 29500 Strict 이름공간과 같은 뜻의 Transitional 이름공간. 관계 유형은 접두어로 바뀐다.
STRICT_MAIN = b"http://purl.oclc.org/ooxml/spreadsheetml/main"
STRICT_NAMESPACES = (
    (STRICT_MAIN, b"http://schemas.openxmlformats.org/spreadsheetml/2006/main"),
    (
        b"http://purl.oclc.org/ooxml/officeDocument/relationships",
        b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    ),
    (
        b"http://purl.oclc.org/ooxml/officeDocument/sharedTypes",
        b"http://schemas.openxmlformats.org/officeDocument/2006/sharedTypes",
    ),
    (
        b"http://purl.oclc.org/ooxml/drawingml/main",
        b"http://schemas.openxmlformats.org/drawingml/2006/main",
    ),
)
# HWPX 본문. 쪽이 아니라 구역마다 파일이 나뉘고 번호는 0부터 이어진다.
SECTION = re.compile(r"Contents/section(\d+)\.xml")
# OWPML 문단 이름공간. 표·행·칸·글자가 모두 여기에 있다.
OWPML = "http://www.hancom.co.kr/hwpml/2011/paragraph"
TABLE = f"{{{OWPML}}}tbl"
ROW = f"{{{OWPML}}}tr"
CELL = f"{{{OWPML}}}tc"
ADDRESS = f"{{{OWPML}}}cellAddr"
CELL_SPAN = f"{{{OWPML}}}cellSpan"
PARAGRAPH = f"{{{OWPML}}}p"
TEXT_TAG = f"{{{OWPML}}}t"
# HTML 표 쪽의 인코딩. 울산 시청·중구·동구 게시판이 모두 UTF-8이다(2026-09-14 실측).
HTML_ENCODING = "utf-8"
# HTML 표식을 찾을 앞부분의 길이. 앞의 공백 줄이 길어도 이 안에 온다.
HTML_WINDOW = 4096
# 화면·표 자체가 원본인 게시판의 표식. HTML에는 고정된 매직 바이트가 없어 서명 대신
# 앞부분의 표식으로 가른다. 2026-09-14 실측: 서울시청·울산시청 상세는 `<!DOCTYPE html>`로,
# 관악 월별 내려받기는 빈 줄 여덟 개 뒤의 `<meta>`로 시작한다.
HTML_MARKERS = (b"<!doctype html", b"<html", b"<table", b"<meta", b"<body")
# 표식을 찾기 전에 걷어낼 앞머리. BOM과 공백만 걷어내고 그 밖의 바이트는 건드리지 않는다.
BOMS = (b"\xef\xbb\xbf", b"\xff\xfe", b"\xfe\xff")
# SpreadsheetML 통합문서의 이름공간. 수집은 받아들이지만(`boards.CONTAINERS`) 격자로는 읽지 않는다.
SPREADSHEETML = b"urn:schemas-microsoft-com:office:spreadsheet"
# 좁은 화면에서만 보이도록 칸마다 되풀이한 열 이름. 값이 아니다(동구 구청장 상세 실측
# `<span class="add-head">금액(원)</span><span class="tds">140,000</span>`).
HTML_REPEATED_LABELS = frozenset({"add-head"})
# 글자가 표의 값이 아닌 요소. 스크립트·스타일 본문은 칸 글자로 옮기지 않는다.
HTML_SKIPPED = frozenset({"script", "style"})
# 닫는 태그가 없는 요소(HTML 표준의 void 요소 가운데 게시판 쪽에 나오는 것).
HTML_VOID = frozenset({"br", "img", "input", "hr", "meta", "link", "col", "wbr", "area", "source"})
# 표 앞에서 이름표로 삼을 문단. 시청 상세의 `<h2>`, 동구 상세의 `<p>`가 그 날의 제목이다.
HTML_LABEL_TAGS = frozenset({"h1", "h2", "h3", "h4", "h5", "h6", "p"})


class UnsupportedFormat(Exception):
    """이번 구현이 읽지 않는 형식(HWP·ZIP 묶음 등). 자동으로 LLM에 넘기지 않는다."""


class UnreadableOriginal(Exception):
    """형식은 맞지만 내용을 격자로 옮기지 못했다."""


@dataclass(frozen=True)
class Span:
    """세로 병합이 덮은 이어짐 칸 하나. 값은 `holder` 행의 같은 열에 있다.

    행은 1부터, 열은 0부터 센다. 병합을 읽지 못한 원본은 병합 없이 읽으며, 그때의 결과는
    병합을 몰랐을 때와 같다.
    """

    row: int
    column: int
    holder: int


@dataclass(frozen=True)
class Table:
    """표 하나. `name`은 산출물의 위치 표기, `label`은 원본이 이 표를 부르는 이름이다.

    이름표는 형식마다 다른 자리에서 온다 — 통합문서는 시트 이름, PDF는 나온 쪽,
    HWPX는 표 바로 앞의 제목 문단, HTML은 `<caption>`(없으면 표 앞의 제목 문단)이다.
    """

    name: str
    label: str
    rows: tuple[tuple[Cell, ...], ...]
    # 세로 병합. `rows`에는 이어짐 칸이 빈 값 그대로 들어 있어 헤더 서명이 원본과 같다.
    spans: tuple[Span, ...] = ()

    def cell(self, row: int, column: int) -> Cell:
        """격자에 적힌 그대로. 행은 1부터, 열은 0부터 센다. 없는 칸은 빈 문자열이다."""
        cells = self.rows[row - 1] if 0 < row <= len(self.rows) else ()
        return cells[column] if column < len(cells) else ""

    def value(self, row: int, column: int) -> Cell:
        """원본에 보이는 값. 세로 병합의 이어짐 칸은 병합이 담은 값이고 나머지는 `cell`과 같다."""
        for span in self.spans:
            if span.row == row and span.column == column:
                return self.cell(span.holder, column)
        return self.cell(row, column)


def text(value: Cell) -> str:
    """셀을 사람이 읽는 한 줄 문자열로 옮긴다. 줄바꿈·연속 공백은 공백 하나로 모은다."""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return " ".join(value.split())


def is_html(content: bytes) -> bool:
    """HTML 문서인지. BOM·공백을 걷어낸 앞부분이 `<`로 시작하고 실측한 표식을 담아야 한다.

    수집도 이 판정으로 화면·표 원본을 가른다(`boards.container_of`). 앞부분만 본다.
    """
    head = content[:HTML_WINDOW]
    for bom in BOMS:
        head = head.removeprefix(bom)
    head = head.lstrip().lower()
    return head.startswith(b"<") and any(marker in head for marker in HTML_MARKERS)


def read_tables(path: Path) -> tuple[Table, ...]:
    """통합문서는 시트마다, PDF는 괘선으로 나뉜 표마다, HWPX는 `<hp:tbl>`마다, HTML 쪽은
    `<table>`마다 표 하나다.

    내용이 없는 시트·표는 표가 아니며 뒤쪽의 빈 칸·빈 행은 잘라 낸다. 위치 표기는
    통합문서가 `sheet1`, PDF·HWPX·HTML이 `table1`이다.
    """
    content = path.read_bytes()
    if content.startswith(OLE2):
        prefix, blocks = "sheet", _legacy(content)
    elif content.startswith(ZIP):
        prefix, blocks = _zip(content)
    elif content.startswith(PDF):
        prefix, blocks = "table", _pdf(content)
    elif SPREADSHEETML in content[:HTML_WINDOW]:
        # SpreadsheetML도 `<Table>`을 담아 HTML 표식에 걸린다. 읽지 않는 형식으로 그대로 둔다.
        raise UnsupportedFormat("SpreadsheetML workbook")
    elif is_html(content):
        prefix, blocks = "table", _html(content)
    else:
        raise UnsupportedFormat("not a workbook, HWPX, PDF, or HTML table page")
    tables = []
    for index, (label, rows, spans) in enumerate(blocks, start=1):
        trimmed = _trim(rows)
        if trimmed:
            tables.append(Table(f"{prefix}{index}", label, trimmed, spans))
    return tuple(tables)


def _trim(rows: list[tuple[Cell, ...]]) -> tuple[tuple[Cell, ...], ...]:
    """행 번호가 원본과 같도록 앞쪽 빈 행은 두고 뒤쪽만 자른다."""
    cut = [cells[: max((i + 1 for i, value in enumerate(cells) if text(value)), default=0)]
           for cells in rows]  # fmt: skip
    while cut and not cut[-1]:
        cut.pop()
    return tuple(cut)


def _legacy(content: bytes) -> list[Block]:
    try:
        # 병합은 서식 기록에 들어 있어 그것까지 읽어야 이어짐 칸을 알 수 있다.
        book = xlrd.open_workbook(file_contents=content, formatting_info=True)
    except xlrd.XLRDError:
        # HWP 5.0도 같은 OLE2 컨테이너를 쓴다. 엑셀 통합문서가 없으면 지원하지 않는 형식이다.
        raise UnsupportedFormat("OLE2 container without an Excel workbook") from None
    except Exception:
        raise UnreadableOriginal("workbook could not be read") from None
    return [
        (
            sheet.name,
            [
                tuple(_cell(sheet, row, column, book.datemode) for column in range(sheet.ncols))
                for row in range(sheet.nrows)
            ],
            _legacy_spans(sheet),
        )
        for sheet in book.sheets()
    ]


def _legacy_spans(sheet: xlrd.sheet.Sheet) -> tuple[Span, ...]:
    """병합 범위는 반쯤 열린 구간이다. `(16, 18, 1, 2)`는 17~18행의 둘째 열 하나를 덮는다."""
    return tuple(
        Span(row + 1, column, top + 1)
        for top, bottom, left, right in sheet.merged_cells
        for row in range(top + 1, bottom)
        for column in range(left, right)
    )


def _zip(content: bytes) -> tuple[str, list[Block]]:
    """ZIP 컨테이너는 엑셀 통합문서이거나 HWPX 본문이다. 둘 다 아니면 읽지 않는다."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = archive.namelist()
    except zipfile.BadZipFile:
        raise UnsupportedFormat("broken ZIP container") from None
    if "xl/workbook.xml" in names:
        return "sheet", _xlsx(content)
    sections = sorted(
        (int(found.group(1)), name)
        for name in names
        if (found := SECTION.fullmatch(name)) is not None
    )
    if not sections:
        # 게시판의 첨부 묶음 등 엑셀도 HWPX도 아닌 ZIP 컨테이너.
        raise UnsupportedFormat("ZIP container without an Excel workbook or an HWPX body")
    return "table", _hwpx(content, [name for _, name in sections])


def _xlsx(content: bytes) -> list[Block]:
    data = _transitional(content)
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise UnreadableOriginal("workbook could not be read") from None
    if not book.worksheets:
        # 시트를 하나도 못 읽은 통합문서를 빈 원본으로 보고하지 않는다.
        raise UnreadableOriginal("workbook has no readable worksheet")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            spans = _xlsx_spans(archive)
        blocks: list[Block] = []
        for sheet in book.worksheets:
            rows: list[tuple[Cell, ...]] = []
            blank_run = 0
            for values in sheet.iter_rows(values_only=True):
                cells = tuple(_xlsx_cell(value) for value in values)
                blank_run = 0 if any(text(value) for value in cells) else blank_run + 1
                # 서식만 남은 거대한 범위를 끝까지 따라가지 않는다.
                if blank_run > MAX_BLANK_RUN:
                    break
                rows.append(cells)
            blocks.append((sheet.title, rows, spans.get(sheet.title, ())))
        book.close()
    except Exception:
        raise UnreadableOriginal("workbook could not be read") from None
    return blocks


def _xlsx_spans(archive: zipfile.ZipFile) -> dict[str, tuple[Span, ...]]:
    """시트 이름마다 세로 병합. openpyxl은 읽기 전용으로 열면 병합을 싣지 않는다.

    openpyxl이 이미 연 통합문서만 여기로 온다. 관계 기록과 `xl/workbook.xml`은 openpyxl도
    같은 것을 읽으므로 둘이 없거나 깨진 통합문서는 여기 오기 전에 읽기 실패다(2026-09-13 실측).
    시트가 아닌 부분을 가리키는 관계만 병합 없이 지나간다.
    """
    names = set(archive.namelist())
    targets = {
        item.get("Id"): item.get("Target", "")
        for item in ElementTree.fromstring(archive.read(WORKBOOK_RELS))
    }
    book = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    spans = {}
    for sheet in book.iterfind(f"{MAIN}sheets/{MAIN}sheet"):
        target = targets.get(sheet.get(f"{DOCUMENT}id", ""), "")
        part = target[1:] if target.startswith("/") else f"xl/{target}"
        spans[sheet.get("name", "")] = _sheet_spans(archive, part) if part in names else ()
    return spans


def _sheet_spans(archive: zipfile.ZipFile, part: str) -> tuple[Span, ...]:
    """시트 XML의 `mergeCell`만 훑는다. 칸 값은 openpyxl이 읽으므로 여기서는 버린다."""
    spans: list[Span] = []
    with archive.open(part) as stream:
        for _, element in ElementTree.iterparse(stream):
            if element.tag == f"{MAIN}mergeCell":
                spans.extend(_merged(element.get("ref", "")))
            element.clear()
    return tuple(spans)


def _merged(reference: str) -> Iterator[Span]:
    """`B17:B18`처럼 적힌 병합 범위. 두 행 이상을 덮을 때만 이어짐 칸이 생긴다."""
    found = re.fullmatch(r"([A-Z]+)([0-9]+):([A-Z]+)([0-9]+)", reference)
    if found is None:
        return
    first, top, last, bottom = found.groups()
    for row in range(int(top) + 1, int(bottom) + 1):
        for column in range(_column(first), _column(last) + 1):
            yield Span(row, column, int(top))


def _column(letters: str) -> int:
    """`A`가 0인 열 번호."""
    index = 0
    for letter in letters:
        index = index * 26 + ord(letter) - ord("A") + 1
    return index - 1


def _pdf(content: bytes) -> list[Block]:
    """괘선으로 칸이 나뉜 표를 격자로 옮긴다. 셀 값은 모두 글자이며 숫자로 바꾸지 않는다.

    글자 층이 없는 스캔본이나 괘선 없는 안내문은 표가 나오지 않는다. 그때는 표 0개로
    돌려 미해결로 남기며, 집행 없음으로 바꾸지 않는다.
    """
    try:
        with pdfplumber.open(io.BytesIO(content)) as document:
            return [
                (f"{page.page_number}쪽", _pdf_rows(table), _pdf_spans(table))
                for page in document.pages
                for table in page.find_tables()
            ]
    except Exception:
        # 암호가 걸렸거나 구조가 깨진 PDF. 원본 내용은 사유에 담지 않는다.
        raise UnreadableOriginal("PDF could not be read") from None


def _pdf_rows(table: RuledTable) -> list[tuple[Cell, ...]]:
    """칸이 없는 자리는 빈 글자로 둔다. 병합 여부는 `_pdf_spans`가 따로 싣는다."""
    return [tuple("" if cell is None else cell for cell in row) for row in table.extract()]


def _pdf_spans(table: RuledTable) -> tuple[Span, ...]:
    """칸이 없는 자리 가운데 위 칸이 그 행 끝까지 내려온 것만 세로 병합이다.

    통합문서와 달리 PDF는 병합된 자리에 칸 자체가 없다. 위 칸이 거기까지 내려오지 않았다면
    그냥 칸이 없는 자리이며 병합으로 보지 않는다.
    """
    spans: list[Span] = []
    # 열마다 값을 담은 마지막 칸의 행과 그 칸의 아랫변. 칸도 행도 (왼쪽, 위, 오른쪽, 아래)다.
    holders: dict[int, tuple[int, float]] = {}
    for row, line in enumerate(table.rows, start=1):
        bottom = line.bbox[3]
        for column, cell in enumerate(line.cells):
            if cell is not None:
                holders[column] = (row, cell[3])
                continue
            holder = holders.get(column)
            if holder is None:
                continue
            holder_row, holder_bottom = holder
            if holder_bottom < bottom - RULE_TOLERANCE:
                # 위 칸이 이 행 끝까지 내려오지 않았다. 병합이 아니라 그냥 칸이 없는 자리다.
                holders.pop(column, None)
                continue
            spans.append(Span(row, column, holder_row))
    return tuple(spans)


@dataclass
class _Body:
    """본문 하나를 훑는 동안의 상태. 표 바로 앞의 문단을 그 표의 이름표로 쓴다."""

    # 표 앞에 문단이 없을 때 쓸 본문 이름(`section0`).
    fallback_label: str
    blocks: list[Block] = field(default_factory=list)
    label: str = ""


def _hwpx(content: bytes, sections: list[str]) -> list[Block]:
    """OWPML 본문의 `<hp:tbl>`마다 표 하나다. 표 밖 문단은 표로 만들지 않는다.

    HWPX에는 시트 이름이 없어 표 바로 앞의 비어 있지 않은 문단을 `label`로 남긴다. 실제
    원본에서 그 자리는 `2026. 1분기 업무추진비 집행내역(회계과)` 같은 표 제목이다.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            bodies = [(name, archive.read(name)) for name in sections]
    except (zipfile.BadZipFile, KeyError):
        raise UnreadableOriginal("HWPX body could not be read") from None
    blocks: list[Block] = []
    for name, body in bodies:
        found = _Body(Path(name).stem)
        try:
            _scan(ElementTree.fromstring(body), found)
        except UnreadableOriginal:
            # 자리 표기를 두고 이미 정한 사유가 있다. 뭉뚱그리지 않고 그대로 올린다.
            raise
        except Exception:
            # 구조가 깨진 본문. 원본 내용은 사유에 담지 않는다.
            raise UnreadableOriginal("HWPX body could not be read") from None
        blocks.extend(found.blocks)
    return blocks


def _scan(element: ElementTree.Element, body: _Body) -> None:
    """본문을 적힌 순서대로 훑는다. 칸 안에 든 표도 제 표로 따로 낸다."""
    for child in element:
        if child.tag == TABLE:
            rows, spans = _hwpx_grid(child)
            body.blocks.append((body.label or body.fallback_label, rows, spans))
            body.label = ""
            for row in child.findall(ROW):
                for cell in row.findall(CELL):
                    _scan(cell, body)
            # 칸 안의 문단은 이 표의 것이다. 다음 표의 이름표로 새지 않게 다시 비운다.
            body.label = ""
            continue
        if child.tag == PARAGRAPH:
            found = _joined_text(child)
            if found.strip():
                body.label = found.strip()
        _scan(child, body)


def _hwpx_grid(table: ElementTree.Element) -> tuple[list[tuple[Cell, ...]], tuple[Span, ...]]:
    """`<hp:cellAddr>`가 가리키는 자리에 칸을 놓아 행·열 번호를 원본과 같게 둔다.

    병합으로 덮인 자리는 비워 두고 세로 병합만 `Span`으로 따로 싣는다. 통합문서의 병합 칸도
    왼쪽 위에만 값이 있고 나머지는 비어 있으므로 `Table.cell`·`Table.value`를 읽는 쪽이
    형식마다 다른 규칙을 알 필요가 없다.
    """
    placed: dict[tuple[int, int], Cell] = {}
    taken: set[tuple[int, int]] = set()
    spans: list[Span] = []
    height = 0
    width = 0
    for index, element in enumerate(table.findall(ROW)):
        column = 0
        for cell in element.findall(CELL):
            while (index, column) in taken:
                column += 1
            row, column = _at(cell, index, column)
            rows, columns = _cell_span(cell)
            height = max(height, row + rows)
            width = max(width, column + columns)
            if height * width > MAX_TABLE_CELLS:
                # 실제 원본의 가장 큰 표도 600칸이 되지 않는다. 여기까지 오면 표기가 깨진
                # 것이며, 자리를 짐작해 줄이지 않고 옮기지 못했다고 남긴다.
                raise UnreadableOriginal("HWPX table is too large to lay out")
            if (row, column) in taken:
                # 두 칸이 한 자리를 가리킨다. 나중 칸으로 덮어써 앞 칸을 조용히 버리지 않는다.
                raise UnreadableOriginal("HWPX cells overlap in the same position")
            placed[(row, column)] = _joined_text(cell)
            taken.update((row + r, column + c) for r in range(rows) for c in range(columns))
            # 세로 병합이 덮은 자리. `Span`의 행은 1부터 세므로 이 칸의 행은 `row + 1`이다.
            spans.extend(
                Span(row + r + 1, column + c, row + 1)
                for r in range(1, rows)
                for c in range(columns)
            )
            column += columns
    grid = [
        tuple(placed.get((row, column), "") for column in range(width)) for row in range(height)
    ]
    return grid, tuple(spans)


def _at(cell: ElementTree.Element, row: int, column: int) -> tuple[int, int]:
    """칸이 놓인 자리. `<hp:cellAddr>`가 없는 표는 적힌 차례대로 왼쪽부터 채운다."""
    address = cell.find(ADDRESS)
    if address is None:
        return row, column
    return _number(address.get("rowAddr"), row), _number(address.get("colAddr"), column)


def _cell_span(cell: ElementTree.Element) -> tuple[int, int]:
    """칸이 덮는 행·열 수. 병합하지 않은 칸은 하나씩이다."""
    span = cell.find(CELL_SPAN)
    if span is None:
        return 1, 1
    return _number(span.get("rowSpan"), 1, 1), _number(span.get("colSpan"), 1, 1)


def _number(value: str | None, fallback: int, floor: int = 0) -> int:
    """자리도 크기도 음수일 수 없다. 읽지 못한 값은 적힌 차례에서 얻은 값으로 둔다."""
    try:
        found = int(value) if value is not None else fallback
    except ValueError:
        found = fallback
    return max(found, floor)


def _joined_text(element: ElementTree.Element) -> str:
    """칸·문단의 글자를 잇는다. 한 칸이 `<hp:t>` 여러 개로 쪼개져 있어도 한 값이다.

    실제 원본은 칸 안에서 줄을 나눠 `결제`·`방법`을 따로 적는다. 그 사이에 무엇도 끼우지
    않아야 원본이 보여 주는 `결제방법`이 된다. 칸 안에 든 표는 제 표로 따로 내므로 여기서
    글자를 가져오지 않는다.
    """
    return "".join(
        "".join(child.itertext()) if child.tag == TEXT_TAG else _joined_text(child)
        for child in element
        if child.tag != TABLE
    )


def _html(content: bytes) -> list[Block]:
    """게시판 쪽의 `<table>`마다 표 하나다. 표 밖 글자는 표로 만들지 않는다.

    이름표는 `<caption>`이고, 없으면 표 바로 앞의 비어 있지 않은 제목·문단이다. 칸이 여러 열·행을
    덮으면(`colspan`·`rowspan`) 값은 왼쪽 위 칸에만 두고 세로로 덮은 자리는 `Span`으로 싣는다 —
    통합문서·HWPX의 병합과 같은 규칙이다. 칸 안에 든 표는 제 표로 따로 낸다.
    """
    try:
        page = content.decode(HTML_ENCODING)
    except UnicodeDecodeError:
        raise UnreadableOriginal("HTML page is not in the measured encoding") from None
    document = _HtmlTables()
    try:
        document.feed(page)
        document.close()
    except UnreadableOriginal:
        raise
    except Exception:
        raise UnreadableOriginal("HTML page could not be read") from None
    return [table.block() for table in document.tables]


@dataclass
class _HtmlCell:
    parts: list[str]
    rows: int
    columns: int


@dataclass
class _HtmlTable:
    label: str
    caption: list[str] | None = None
    rows: list[list[_HtmlCell]] = field(default_factory=list)

    def block(self) -> Block:
        placed: dict[tuple[int, int], Cell] = {}
        taken: set[tuple[int, int]] = set()
        spans: list[Span] = []
        width = 0
        for index, cells in enumerate(self.rows):
            column = 0
            for cell in cells:
                while (index, column) in taken:
                    column += 1
                if len(self.rows) * max(width, column + cell.columns) > MAX_TABLE_CELLS:
                    # 병합 표기가 깨진 쪽이 격자를 키우지 못하게 한다(HWPX와 같은 상한).
                    raise UnreadableOriginal("HTML table is too large to lay out")
                placed[(index, column)] = " ".join(" ".join(cell.parts).split())
                # 표 끝을 넘는 세로 병합은 없는 행을 덮으므로 표 안의 행까지만 차지한다.
                height = min(cell.rows, len(self.rows) - index)
                taken.update(
                    (index + r, column + c) for r in range(height) for c in range(cell.columns)
                )
                spans.extend(
                    Span(index + r + 1, column + c, index + 1)
                    for r in range(1, height)
                    for c in range(cell.columns)
                )
                column += cell.columns
                width = max(width, column)
        rows = [
            tuple(placed.get((row, column), "") for column in range(width))
            for row in range(len(self.rows))
        ]
        caption = " ".join(" ".join(self.caption or ()).split())
        return caption or self.label, rows, tuple(spans)


class _HtmlTables(HTMLParser):
    """표·행·칸과 표 앞의 제목 문단만 따라간다. 요소 구조의 다른 부분에는 기대지 않는다."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[_HtmlTable] = []
        self._open: list[_HtmlTable] = []
        self._cell: list[_HtmlCell | None] = []
        # 글자를 버리는 요소들. 되풀이한 열 이름·스크립트 안이면 비어 있지 않다.
        self._hidden: list[str] = []
        # 지금 글자를 모으는 캡션·제목 문단. 닫히면 비운다.
        self._caption: list[str] | None = None
        self._heading: list[str] | None = None
        self._label = ""

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        values = dict(attrs)
        if tag in HTML_VOID:
            # 닫는 태그가 없는 요소는 숨길 글자도 없다. 숨긴 요소 목록에 넣으면 영영 닫히지 않는다.
            self.handle_startendtag(tag, attrs)
            return
        if self._hidden:
            self._hidden.append(tag)
            return
        if tag in HTML_SKIPPED or HTML_REPEATED_LABELS & set((values.get("class") or "").split()):
            self._hidden.append(tag)
            return
        if tag == "table":
            table = _HtmlTable(self._label)
            self.tables.append(table)
            self._open.append(table)
            self._cell.append(None)
            self._label = ""
            return
        if not self._open:
            if tag in HTML_LABEL_TAGS:
                self._heading = []
            return
        table = self._open[-1]
        if tag == "caption":
            table.caption = self._caption = []
        elif tag == "tr":
            table.rows.append([])
        elif tag in {"td", "th"}:
            if not table.rows:
                table.rows.append([])
            cell = _HtmlCell([], _span(values.get("rowspan")), _span(values.get("colspan")))
            table.rows[-1].append(cell)
            self._cell[-1] = cell

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        # `<br/>`처럼 닫는 태그가 없는 요소는 여는 태그만 본다. 숨긴 깊이를 늘리지 않는다.
        # 열 이름 표시가 붙은 `<br>`도 칸의 줄바꿈일 뿐 숨길 글자가 없다.
        if tag == "br" and not self._hidden and self._open and self._cell[-1] is not None:
            self._cell[-1].parts.append(" ")

    def handle_endtag(self, tag: str) -> None:
        if self._hidden:
            # 닫는 태그를 빠뜨린 쪽도 있어 같은 이름이 나올 때까지 거슬러 닫는다.
            while self._hidden and self._hidden.pop() != tag:
                pass
            return
        if tag == "table" and self._open:
            self._open.pop()
            self._cell.pop()
            return
        if not self._open:
            if tag in HTML_LABEL_TAGS and self._heading is not None:
                found = " ".join(" ".join(self._heading).split())
                if found:
                    self._label = found
                self._heading = None
            return
        if tag in {"td", "th"}:
            self._cell[-1] = None
        elif tag == "caption":
            self._caption = None

    def handle_data(self, data: str) -> None:
        if self._hidden:
            return
        if not self._open:
            if self._heading is not None:
                self._heading.append(data)
            return
        cell = self._cell[-1]
        if cell is not None:
            cell.parts.append(data)
        elif self._caption is not None:
            self._caption.append(data)


def _span(value: str | None) -> int:
    """칸이 덮는 행·열 수. 읽지 못한 값은 칸 하나다."""
    try:
        return max(int(value or "1"), 1)
    except ValueError:
        return 1


def _transitional(content: bytes) -> bytes:
    """ISO Strict 이름공간을 Transitional로 바꾼다. openpyxl은 Strict 시트를 조용히 건너뛴다."""
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if STRICT_MAIN not in archive.read("xl/workbook.xml"):
            return content
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as target:
            for item in archive.infolist():
                data = archive.read(item)
                if item.filename.endswith((".xml", ".rels")):
                    for strict, transitional in STRICT_NAMESPACES:
                        data = data.replace(strict, transitional)
                target.writestr(item, data)
    return stream.getvalue()


def _xlsx_cell(value: object) -> Cell:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int | float):
        return float(value)
    return str(value).strip()


def _cell(sheet: xlrd.sheet.Sheet, row: int, column: int, datemode: int) -> Cell:
    kind = sheet.cell_type(row, column)
    value = sheet.cell_value(row, column)
    if kind == xlrd.XL_CELL_DATE:
        try:
            return _as_datetime(xlrd.xldate_as_datetime(value, datemode))
        except (ValueError, OverflowError, xlrd.xldate.XLDateError):
            raise UnreadableOriginal("invalid date cell") from None
    if kind == xlrd.XL_CELL_NUMBER:
        return float(value)
    if kind in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return ""
    return str(value).strip()


def _as_datetime(value: object) -> datetime:
    if not isinstance(value, datetime):
        raise UnreadableOriginal("invalid date cell")
    return value
